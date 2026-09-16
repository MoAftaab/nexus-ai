"""Hackathon Excel workbook ingestion service.

Parses the Warehouse_AI_Hackathon_Synthetic_Dataset workbook and normalises
all six SAP-style data sheets into typed Python dicts.
"""
from __future__ import annotations

import logging
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl

logger = logging.getLogger(__name__)

REQUIRED_SHEETS = {
    "Material_Master",
    "Inventory_Stock",
    "Warehouse_Bin",
    "Deliveries_Dispatch",
    "Purchase_Replenish",
    "Vendor_Master",
}

_EXCEL_EPOCH = datetime(1899, 12, 30, tzinfo=timezone.utc)
_NULL_SENTINELS = {"", "n/a", "na", "#n/a", "#null!", "#ref!", "none", "null"}


def _clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        if stripped.lower() in _NULL_SENTINELS:
            return None
        return stripped or None
    return value


def _parse_date(value: Any) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        s = value.strip()
        if not s or s.lower() in _NULL_SENTINELS:
            return None
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                pass
        return None
    if isinstance(value, (int, float)):
        import datetime as dt_mod
        try:
            return (_EXCEL_EPOCH + dt_mod.timedelta(days=int(value))).date()
        except Exception:
            return None
    return None


def _parse_float(value: Any) -> float | None:
    c = _clean(value)
    if c is None:
        return None
    if isinstance(c, str):
        c = c.replace(",", "").replace(" ", "")
    try:
        return float(c)
    except Exception:
        return None


def _str_upper(v: Any) -> str | None:
    c = _clean(v)
    return c.upper() if isinstance(c, str) else None


def _str_exact(v: Any) -> str | None:
    return _clean(v)


def _normalize_material_master(row: tuple) -> dict[str, Any]:
    p = (*row, *([None] * 12))
    m = p[:12]
    return {
        "material": _str_exact(m[0]),
        "description": _str_exact(m[1]),
        "material_type": _str_upper(m[2]),
        "material_group": _str_exact(m[3]),
        "base_uom": _str_exact(m[4]),
        "plant": _str_exact(m[5]),
        "reorder_point": _parse_float(m[6]),
        "safety_stock": _parse_float(m[7]),
        "lead_time_days": _parse_float(m[8]),
        "abc_class": _str_upper(m[9]),
        "hazmat_flag": _str_upper(m[10]),
        "lifecycle_status": _str_upper(m[11]),
    }


def _normalize_inventory_stock(row: tuple) -> dict[str, Any]:
    p = (*row, *([None] * 10))
    m = p[:10]
    return {
        "material": _str_exact(m[0]),
        "plant": _str_exact(m[1]),
        "storage_location": _str_exact(m[2]),
        "batch": _str_exact(m[3]),
        "uom": _str_exact(m[4]),
        "qty_on_hand": _parse_float(m[5]),
        "blocked_qty": _parse_float(m[6]),
        "in_transit_qty": _parse_float(m[7]),
        "batch_expiry": _parse_date(m[8]),
        "last_movement_date": _parse_date(m[9]),
    }


def _normalize_warehouse_bin(row: tuple) -> dict[str, Any]:
    p = (*row, *([None] * 7))
    m = p[:7]
    return {
        "bin": _str_exact(m[0]),
        "storage_type": _str_upper(m[1]),
        "assigned_material": _str_exact(m[2]),
        "capacity": _parse_float(m[3]),
        "occupied": _parse_float(m[4]),
        "bin_status": _str_upper(m[5]),
        "plant": _str_exact(m[6]),
    }


def _normalize_deliveries_dispatch(row: tuple) -> dict[str, Any]:
    p = (*row, *([None] * 9))
    m = p[:9]
    return {
        "delivery": _str_exact(m[0]),
        "material": _str_exact(m[1]),
        "plant": _str_exact(m[2]),
        "order_qty": _parse_float(m[3]),
        "ship_to": _str_exact(m[4]),
        "route": _str_exact(m[5]),
        "created_date": _parse_date(m[6]),
        "planned_gi_date": _parse_date(m[7]),
        "status": _str_upper(m[8]),
    }


def _normalize_purchase_replenish(row: tuple) -> dict[str, Any]:
    p = (*row, *([None] * 10))
    m = p[:10]
    return {
        "purchase_order": _str_exact(m[0]),
        "material": _str_exact(m[1]),
        "vendor": _str_exact(m[2]),
        "plant": _str_exact(m[3]),
        "po_qty": _parse_float(m[4]),
        "unit_price": _parse_float(m[5]),
        "currency": _str_exact(m[6]),
        "order_date": _parse_date(m[7]),
        "expected_delivery": _parse_date(m[8]),
        "po_status": _str_upper(m[9]),
    }


def _normalize_vendor_master(row: tuple) -> dict[str, Any]:
    p = (*row, *([None] * 6))
    m = p[:6]
    return {
        "vendor": _str_exact(m[0]),
        "vendor_name": _str_exact(m[1]),
        "country": _str_exact(m[2]),
        "quality_rating": _str_upper(m[3]),
        "otd_pct": _parse_float(m[4]),
        "procurement_block": _str_upper(m[5]),
    }


_NORMALIZERS = {
    "Material_Master": _normalize_material_master,
    "Inventory_Stock": _normalize_inventory_stock,
    "Warehouse_Bin": _normalize_warehouse_bin,
    "Deliveries_Dispatch": _normalize_deliveries_dispatch,
    "Purchase_Replenish": _normalize_purchase_replenish,
    "Vendor_Master": _normalize_vendor_master,
}


def _sheet_to_dicts(sheet, normalizer) -> list[dict[str, Any]]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    result = []
    for rr in rows[1:]:
        if not any(rr):
            continue
        try:
            result.append(normalizer(rr))
        except Exception as exc:
            logger.warning("Row skipped: %s", exc)
    return result


def load_hackathon_workbook(path: Path | str) -> dict[str, Any]:
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    available = set(wb.sheetnames)
    missing_sheets = sorted(REQUIRED_SHEETS - available)
    extra_sheets = sorted(available - REQUIRED_SHEETS)
    if missing_sheets:
        logger.warning("Missing required sheets: %s", missing_sheets)
    result: dict[str, Any] = {"missing_sheets": missing_sheets, "extra_sheets": extra_sheets}
    key_map = {
        "Material_Master": "material_master",
        "Inventory_Stock": "inventory_stock",
        "Warehouse_Bin": "warehouse_bin",
        "Deliveries_Dispatch": "deliveries_dispatch",
        "Purchase_Replenish": "purchase_replenish",
        "Vendor_Master": "vendor_master",
    }
    for sn, rk in key_map.items():
        if sn in available:
            recs = _sheet_to_dicts(wb[sn], _NORMALIZERS[sn])
            result[rk] = recs
            logger.info("Loaded %s: %d rows", sn, len(recs))
        else:
            result[rk] = []
    wb.close()
    return result


def build_material_index(mm: list[dict[str, Any]]) -> dict[tuple[str, str], dict[str, Any]]:
    return {(r["material"], r["plant"]): r for r in mm if r.get("material") and r.get("plant")}


def build_material_set(mm: list[dict[str, Any]]) -> set[str]:
    return {r["material"] for r in mm if r.get("material")}


def build_vendor_index(vm: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {r["vendor"]: r for r in vm if r.get("vendor")}


def build_inventory_by_material_plant(inv: list[dict[str, Any]]) -> dict[tuple[str, str], list[dict[str, Any]]]:
    result: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for row in inv:
        key = (row.get("material", ""), row.get("plant", ""))
        result.setdefault(key, []).append(row)
    return result
